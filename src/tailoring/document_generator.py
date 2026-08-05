import os
import subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable

class DocumentGenerator:
    """Generates exact LaTeX PDF compilations, Word documents, and Cover Letter PDFs."""

    @staticmethod
    def compile_latex_to_pdf(tex_path: str, output_pdf_path: str) -> bool:
        """
        Compiles a .tex file directly to a PDF using the local Tectonic engine,
        preserving 100% exact colors, layout, fonts, margins, and LaTeX macros.
        """
        tex_dir = os.path.dirname(os.path.abspath(tex_path))
        tex_filename = os.path.basename(tex_path)
        tectonic_bin = os.path.abspath("bin/tectonic") if os.path.exists("bin/tectonic") else os.path.abspath("tectonic")

        cmd = [tectonic_bin, tex_filename] if os.path.exists(tectonic_bin) else ["tectonic", tex_filename]

        try:
            res = subprocess.run(cmd, cwd=tex_dir, capture_output=True, text=True, check=True)
            default_pdf = os.path.join(tex_dir, tex_filename.replace(".tex", ".pdf"))
            if os.path.exists(default_pdf) and default_pdf != output_pdf_path:
                os.rename(default_pdf, output_pdf_path)
            print(f"   ✅ Successfully compiled exact LaTeX PDF: {output_pdf_path}")
            return True
        except Exception as e:
            print(f"   ⚠️ Tectonic compilation error for {tex_path}: {e}")
            return False

    @staticmethod
    def create_cover_letter_docx(letter_data: dict, output_path: str):
        """Creates a professional cover_letter.docx file."""
        doc = Document()
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        primary_color = RGBColor(11, 92, 140)
        dark_text = RGBColor(17, 17, 17)
        body_text = RGBColor(34, 34, 34)

        p_name = doc.add_paragraph()
        r_name = p_name.add_run(letter_data.get("applicant_name", "MUHAMMAD FAHAD"))
        r_name.font.size = Pt(18)
        r_name.font.bold = True
        r_name.font.color.rgb = primary_color

        p_cont = doc.add_paragraph()
        r_cont = p_cont.add_run(letter_data.get("applicant_contact", "Lahore, Pakistan | devfahad785@gmail.com | +92-3329296026"))
        r_cont.font.size = Pt(10)
        r_cont.font.color.rgb = RGBColor(100, 100, 100)
        p_cont.paragraph_format.space_after = Pt(18)

        p_meta = doc.add_paragraph()
        p_meta.paragraph_format.space_after = Pt(14)
        r_meta = p_meta.add_run(f"Date: {letter_data.get('date', 'August 2026')}\n")
        r_meta.add_text(f"To: {letter_data.get('hiring_manager', 'Hiring Team')}\n")
        r_meta.add_text(f"Company: {letter_data.get('company_name', 'Hiring Company')}\n")
        r_meta.add_text(f"Re: Application for {letter_data.get('job_title', 'Software Engineer')}")
        r_meta.font.size = Pt(10.5)
        r_meta.font.color.rgb = dark_text

        paragraphs = letter_data.get("body_paragraphs", [])
        if isinstance(paragraphs, str):
            paragraphs = [paragraphs]

        for p_text in paragraphs:
            p_body = doc.add_paragraph()
            p_body.paragraph_format.space_after = Pt(10)
            p_body.paragraph_format.line_spacing = 1.15
            r_b = p_body.add_run(p_text)
            r_b.font.size = Pt(11)
            r_b.font.color.rgb = body_text

        p_sign = doc.add_paragraph()
        p_sign.paragraph_format.space_before = Pt(14)
        r_sign = p_sign.add_run("Sincerely,\n\nMuhammad Fahad")
        r_sign.font.size = Pt(11)
        r_sign.font.bold = True
        r_sign.font.color.rgb = primary_color

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        doc.save(output_path)

    @staticmethod
    def create_cover_letter_pdf(letter_data: dict, output_path: str):
        """Creates a professional cover_letter.pdf file using ReportLab."""
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        doc = SimpleDocTemplate(
            output_path,
            pagesize=letter,
            rightMargin=54,
            leftMargin=54,
            topMargin=54,
            bottomMargin=54
        )

        styles = getSampleStyleSheet()
        primary_color = colors.HexColor("#0B5C8C")
        dark_color = colors.HexColor("#111111")
        body_color = colors.HexColor("#222222")

        style_name = ParagraphStyle(
            'NameStyle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=20,
            leading=24,
            textColor=primary_color,
            spaceAfter=4
        )

        style_contact = ParagraphStyle(
            'ContactStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=13,
            textColor=colors.HexColor("#555555"),
            spaceAfter=14
        )

        style_meta = ParagraphStyle(
            'MetaStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10.5,
            leading=15,
            textColor=dark_color,
            spaceAfter=14
        )

        style_body = ParagraphStyle(
            'BodyStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10.5,
            leading=15,
            textColor=body_color,
            spaceAfter=10
        )

        story = []
        name = letter_data.get("applicant_name", "MUHAMMAD FAHAD").upper()
        contact = letter_data.get("applicant_contact", "Lahore, Pakistan | devfahad785@gmail.com | +92-3329296026")
        story.append(Paragraph(name, style_name))
        story.append(Paragraph(contact, style_contact))
        story.append(HRFlowable(width="100%", thickness=1, color=primary_color, spaceAfter=14))

        date_str = letter_data.get('date', 'August 2026')
        manager = letter_data.get('hiring_manager', 'Hiring Manager / Talent Acquisition Team')
        company = letter_data.get('company_name', 'Hiring Company')
        title = letter_data.get('job_title', 'Software Engineer')

        meta_text = f"<b>Date:</b> {date_str}<br/><b>To:</b> {manager}<br/><b>Company:</b> {company}<br/><b>Re:</b> Application for {title}"
        story.append(Paragraph(meta_text, style_meta))
        story.append(Spacer(1, 8))

        paragraphs = letter_data.get("body_paragraphs", [])
        if isinstance(paragraphs, str):
            paragraphs = [paragraphs]

        for p_text in paragraphs:
            clean_p = p_text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(clean_p, style_body))

        story.append(Spacer(1, 14))
        sign_text = "Sincerely,<br/><br/><b>Muhammad Fahad</b>"
        story.append(Paragraph(sign_text, style_body))

        doc.build(story)

    @staticmethod
    def create_cv_docx(cv_data: dict, output_path: str):
        """Creates a professional ATS-friendly tailored_cv.docx file."""
        doc = Document()
        for section in doc.sections:
            section.top_margin = Inches(0.7)
            section.bottom_margin = Inches(0.7)
            section.left_margin = Inches(0.7)
            section.right_margin = Inches(0.7)

        primary_color = RGBColor(11, 92, 140)
        dark_text = RGBColor(17, 17, 17)
        body_text = RGBColor(34, 34, 34)

        # Header
        name = cv_data.get("name", "MUHAMMAD FAHAD")
        p_name = doc.add_paragraph()
        p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_name = p_name.add_run(name.upper())
        r_name.font.size = Pt(20)
        r_name.font.bold = True
        r_name.font.color.rgb = primary_color

        headline = cv_data.get("headline", "")
        if headline:
            p_head = doc.add_paragraph()
            p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_head = p_head.add_run(headline)
            r_head.font.size = Pt(12)
            r_head.font.bold = True
            r_head.font.color.rgb = dark_text

        contact = cv_data.get("contact", "")
        if contact:
            p_cont = doc.add_paragraph()
            p_cont.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_cont = p_cont.add_run(contact)
            r_cont.font.size = Pt(9.5)
            r_cont.font.color.rgb = RGBColor(80, 80, 80)
            p_cont.paragraph_format.space_after = Pt(12)

        def add_section_heading(title: str):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(title.upper())
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = primary_color

        # Summary
        summary = cv_data.get("summary", "")
        if summary:
            add_section_heading("Professional Summary")
            p_sum = doc.add_paragraph()
            p_sum.paragraph_format.space_after = Pt(8)
            p_sum.paragraph_format.line_spacing = 1.15
            r_sum = p_sum.add_run(summary)
            r_sum.font.size = Pt(10.5)
            r_sum.font.color.rgb = body_text

        # Skills
        skills = cv_data.get("skills", {})
        if skills:
            add_section_heading("Technical Skills")
            for category, skill_list in skills.items():
                p_sk = doc.add_paragraph()
                p_sk.paragraph_format.space_after = Pt(3)
                r_cat = p_sk.add_run(f"{category}: ")
                r_cat.font.bold = True
                r_cat.font.size = Pt(10)
                r_cat.font.color.rgb = dark_text
                
                skills_str = ", ".join(skill_list) if isinstance(skill_list, list) else str(skill_list)
                r_val = p_sk.add_run(skills_str)
                r_val.font.size = Pt(10)
                r_val.font.color.rgb = body_text

        # Experience
        experience = cv_data.get("experience", [])
        if experience:
            add_section_heading("Professional Experience")
            for exp in experience:
                p_exp = doc.add_paragraph()
                p_exp.paragraph_format.space_before = Pt(6)
                p_exp.paragraph_format.space_after = Pt(2)
                
                r_title = p_exp.add_run(exp.get("title", ""))
                r_title.font.bold = True
                r_title.font.size = Pt(11)
                r_title.font.color.rgb = dark_text
                
                company_str = exp.get("company", "")
                location_str = exp.get("location", "")
                dates_str = exp.get("dates", "")
                meta_parts = [p for p in [company_str, location_str, dates_str] if p]
                if meta_parts:
                    r_meta = p_exp.add_run(f" | {' • '.join(meta_parts)}")
                    r_meta.font.italic = True
                    r_meta.font.size = Pt(10)
                    r_meta.font.color.rgb = RGBColor(100, 100, 100)

                for bullet in exp.get("bullets", []):
                    p_b = doc.add_paragraph(style='List Bullet')
                    p_b.paragraph_format.space_after = Pt(2)
                    p_b.paragraph_format.line_spacing = 1.15
                    r_b = p_b.add_run(bullet)
                    r_b.font.size = Pt(10)
                    r_b.font.color.rgb = body_text

        # Projects
        projects = cv_data.get("projects", [])
        if projects:
            add_section_heading("Key Projects")
            for proj in projects:
                p_proj = doc.add_paragraph()
                p_proj.paragraph_format.space_before = Pt(6)
                p_proj.paragraph_format.space_after = Pt(2)
                
                r_ptitle = p_proj.add_run(proj.get("title", ""))
                r_ptitle.font.bold = True
                r_ptitle.font.size = Pt(11)
                r_ptitle.font.color.rgb = dark_text
                
                tech_str = proj.get("tech", "")
                if tech_str:
                    r_ptech = p_proj.add_run(f" ({tech_str})")
                    r_ptech.font.italic = True
                    r_ptech.font.size = Pt(10)
                    r_ptech.font.color.rgb = RGBColor(100, 100, 100)

                for bullet in proj.get("bullets", []):
                    p_b = doc.add_paragraph(style='List Bullet')
                    p_b.paragraph_format.space_after = Pt(2)
                    p_b.paragraph_format.line_spacing = 1.15
                    r_b = p_b.add_run(bullet)
                    r_b.font.size = Pt(10)
                    r_b.font.color.rgb = body_text

        # Education
        education = cv_data.get("education", [])
        if education:
            add_section_heading("Education")
            for edu in education:
                p_edu = doc.add_paragraph()
                p_edu.paragraph_format.space_after = Pt(3)
                r_deg = p_edu.add_run(edu.get("degree", ""))
                r_deg.font.bold = True
                r_deg.font.size = Pt(10.5)
                r_deg.font.color.rgb = dark_text
                
                inst = edu.get("institution", "")
                dates = edu.get("dates", "")
                parts = [p for p in [inst, dates] if p]
                if parts:
                    r_inst = p_edu.add_run(f" — {' | '.join(parts)}")
                    r_inst.font.size = Pt(10)
                    r_inst.font.color.rgb = body_text

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        doc.save(output_path)

    @staticmethod
    def create_cv_pdf(cv_data: dict, output_path: str):
        """Creates a professional, ATS-friendly tailored_cv.pdf file using ReportLab."""
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        doc = SimpleDocTemplate(
            output_path,
            pagesize=letter,
            rightMargin=40,
            leftMargin=40,
            topMargin=40,
            bottomMargin=40
        )

        styles = getSampleStyleSheet()
        primary_color = colors.HexColor("#0B5C8C")
        dark_color = colors.HexColor("#111111")
        body_color = colors.HexColor("#222222")

        style_name = ParagraphStyle(
            'CVNameStyle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=18,
            leading=22,
            alignment=1,
            textColor=primary_color,
            spaceAfter=3
        )

        style_headline = ParagraphStyle(
            'CVHeadlineStyle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=14,
            alignment=1,
            textColor=dark_color,
            spaceAfter=3
        )

        style_contact = ParagraphStyle(
            'CVContactStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            alignment=1,
            textColor=colors.HexColor("#555555"),
            spaceAfter=10
        )

        style_section = ParagraphStyle(
            'CVSectionStyle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=14,
            textColor=primary_color,
            spaceBefore=8,
            spaceAfter=4
        )

        style_body = ParagraphStyle(
            'CVBodyStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9.5,
            leading=13,
            textColor=body_color,
            spaceAfter=4
        )

        style_bullet = ParagraphStyle(
            'CVBulletStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9.5,
            leading=13,
            textColor=body_color,
            leftIndent=12,
            spaceAfter=3
        )

        def escape_txt(text: str) -> str:
            if not text:
                return ""
            return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

        story = []

        name = cv_data.get("name", "MUHAMMAD FAHAD").upper()
        headline = cv_data.get("headline", "")
        contact = cv_data.get("contact", "")

        story.append(Paragraph(escape_txt(name), style_name))
        if headline:
            story.append(Paragraph(escape_txt(headline), style_headline))
        if contact:
            story.append(Paragraph(escape_txt(contact), style_contact))
        story.append(HRFlowable(width="100%", thickness=1, color=primary_color, spaceAfter=8))

        # Summary
        summary = cv_data.get("summary", "")
        if summary:
            story.append(Paragraph("PROFESSIONAL SUMMARY", style_section))
            story.append(Paragraph(escape_txt(summary), style_body))

        # Skills
        skills = cv_data.get("skills", {})
        if skills:
            story.append(Paragraph("TECHNICAL SKILLS", style_section))
            for cat, s_list in skills.items():
                s_str = ", ".join(s_list) if isinstance(s_list, list) else str(s_list)
                line = f"<b>{escape_txt(cat)}:</b> {escape_txt(s_str)}"
                story.append(Paragraph(line, style_body))

        # Experience
        experience = cv_data.get("experience", [])
        if experience:
            story.append(Paragraph("PROFESSIONAL EXPERIENCE", style_section))
            for exp in experience:
                title_str = escape_txt(exp.get("title", ""))
                comp_str = escape_txt(exp.get("company", ""))
                dates_str = escape_txt(exp.get("dates", ""))
                loc_str = escape_txt(exp.get("location", ""))
                meta = " • ".join([p for p in [comp_str, loc_str, dates_str] if p])

                heading_text = f"<b>{title_str}</b>"
                if meta:
                    heading_text += f" <font color='#666666'>| <i>{meta}</i></font>"
                story.append(Paragraph(heading_text, style_body))

                for bullet in exp.get("bullets", []):
                    story.append(Paragraph(f"• {escape_txt(bullet)}", style_bullet))

        # Projects
        projects = cv_data.get("projects", [])
        if projects:
            story.append(Paragraph("KEY PROJECTS", style_section))
            for proj in projects:
                p_title = escape_txt(proj.get("title", ""))
                p_tech = escape_txt(proj.get("tech", ""))
                head = f"<b>{p_title}</b>"
                if p_tech:
                    head += f" <font color='#666666'>| <i>{p_tech}</i></font>"
                story.append(Paragraph(head, style_body))
                for bullet in proj.get("bullets", []):
                    story.append(Paragraph(f"• {escape_txt(bullet)}", style_bullet))

        # Education
        education = cv_data.get("education", [])
        if education:
            story.append(Paragraph("EDUCATION", style_section))
            for edu in education:
                deg = escape_txt(edu.get("degree", ""))
                inst = escape_txt(edu.get("institution", ""))
                dates = escape_txt(edu.get("dates", ""))
                line = f"<b>{deg}</b> — {inst}"
                if dates:
                    line += f" ({dates})"
                story.append(Paragraph(line, style_body))

        doc.build(story)


